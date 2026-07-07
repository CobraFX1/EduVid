"""
Update CHAPTER THREE (Updated).docx — fix all inconsistencies found in the audit.
Uses exact text matching to find and replace specific passages.
"""
from docx import Document
import re, copy

DOC_PATH = r"c:\Projects\eduvid\CHAPTER THREE (Updated).docx"

doc = Document(DOC_PATH)

# ──────────────────────────────────────────
# Helper: replace text in a paragraph while preserving formatting
# ──────────────────────────────────────────
def replace_in_paragraph(paragraph, old_text, new_text):
    """Replace old_text with new_text in a paragraph, keeping the first run's formatting."""
    full = paragraph.text
    if old_text not in full:
        return False
    new_full = full.replace(old_text, new_text)
    # Rebuild: put everything into the first run, clear the rest
    if paragraph.runs:
        paragraph.runs[0].text = new_full
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_full
    return True

def replace_across_doc(old_text, new_text):
    """Try to replace in every paragraph. Returns count of replacements."""
    count = 0
    for p in doc.paragraphs:
        if replace_in_paragraph(p, old_text, new_text):
            count += 1
    # Also check table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if replace_in_paragraph(p, old_text, new_text):
                        count += 1
    return count

# ══════════════════════════════════════════
# REPLACEMENTS
# ══════════════════════════════════════════
replacements = {}

# ─── 1. Bootstrap 5 grid system (§3.3.2) ───
replacements["R01-NFR-Responsiveness"] = replace_across_doc(
    "A mobile-first design using Bootstrap 5 to ensure the platform works seamlessly on smartphones, tablets, and PCs.",
    "A responsive interface using custom CSS Grid and Flexbox layouts with Bootstrap 5 utility support, ensuring the platform works seamlessly on smartphones, tablets, and PCs."
)

# ─── 2. Bootstrap 5 in §3.4.1 ───
replacements["R02-Bootstrap-341"] = replace_across_doc(
    "Bootstrap 5: To ensure a mobile-first experience for Dominion University students, Bootstrap 5 provides a responsive grid system. This ensures the platform is accessible across various devices, including smartphones, tablets, and laptops",
    "Custom CSS3 with Bootstrap 5 Utilities: The platform\u2019s responsive design is achieved through a custom CSS3 design system built with CSS Grid and Flexbox layouts, featuring Glassmorphism visual effects via backdrop-filter. Bootstrap 5 is included as a utility library for supplementary spacing and icon support, while all core layout and component styling is handled through purpose-built CSS classes and scoped @media breakpoints to ensure the interface adapts across smartphones, tablets, and laptops"
)

# ─── 3. Bootstrap 5 in §3.7 ───
replacements["R03-Bootstrap-37"] = replace_across_doc(
    "Bootstrap 5 & CSS3 (UI Framework): Provides a responsive grid system and pre-styled components. This ensures the platform adapts to various screen sizes, from smartphones to laptops, while custom CSS3 handles advanced effects like Glassmorphism.",
    "Custom CSS3 with Bootstrap 5 Utilities (UI Styling): The platform\u2019s visual identity is built on a custom CSS3 design system featuring CSS Grid layouts, Flexbox positioning, and Glassmorphism effects achieved via backdrop-filter. Bootstrap 5 is included as a supplementary utility library, with its icon set (Bootstrap Icons) used extensively throughout the interface. All responsive breakpoints are handled through custom @media queries tailored to the application\u2019s specific layout requirements."
)

# ─── 4. YouTube IFrame API (§3.4.3) ───
replacements["R04-IFrame-API"] = replace_across_doc(
    "Frontend receives the IDs and uses the YouTube IFrame API to embed the video players directly into the student\u2019s dashboard.",
    "Frontend receives the IDs and renders embedded YouTube video players directly into the student\u2019s dashboard using standard iframe embeds with YouTube\u2019s /embed/ endpoint."
)
# Also try with straight quotes
if replacements["R04-IFrame-API"] == 0:
    replacements["R04-IFrame-API"] = replace_across_doc(
        "Frontend receives the IDs and uses the YouTube IFrame API to embed the video players directly into the student's dashboard.",
        "Frontend receives the IDs and renders embedded YouTube video players directly into the student's dashboard using standard iframe embeds with YouTube's /embed/ endpoint."
    )

# ─── 5. Cloud Storage (§3.4.2) ───
replacements["R05-CloudStorage-342"] = replace_across_doc(
    "Cloud Storage: Reserved for storing static assets, such as user profile pictures and application icons.",
    "Image Hosting: Profile pictures and user avatars are hosted via a third-party image hosting service (ImgBB), with the resulting URLs stored directly in the Firestore user document."
)

# ─── 6. API Quota Management (§3.4.2) ───
replacements["R06-QuotaMgmt-342"] = replace_across_doc(
    "API Quota Management: Logic to cache responses and restructure requests to stay within the YouTube Data API v3 daily limits.",
    "API Quota Management: Logic to track daily API usage against the YouTube Data API v3 quota limits. The backend implements an in-memory response cache with TTL-based expiration and a quota tracker that monitors daily unit consumption, ensuring that automated validation routines stay within the daily allowance."
)

# ─── 7. Search & Discovery (§3.3.1) ───
replacements["R07-SearchDiscovery"] = replace_across_doc(
    "Search and Discovery Engine: Filtering and search functionality based on course codes and academic keywords.",
    "Search and Discovery Engine: Text-based search across video titles, topics, and course codes, with dropdown filtering by Department, Course Code, and Level."
)

# ─── 8. Admin Governance (§3.3.1) — flagging clarification ───
replacements["R08-AdminGov"] = replace_across_doc(
    'Administrator Governance Dashboard: A specialized interface for university admins to moderate content by "flagging" videos and managing user permissions.',
    "Administrator Governance Dashboard: A specialized interface for university admins to review flagged content, manage user permissions, and moderate inappropriate comments."
)
# Try with smart quotes
if replacements["R08-AdminGov"] == 0:
    replacements["R08-AdminGov"] = replace_across_doc(
        'Administrator Governance Dashboard: A specialized interface for university admins to moderate content by \u201cflagging\u201d videos and managing user permissions.',
        "Administrator Governance Dashboard: A specialized interface for university admins to review flagged content, manage user permissions, and moderate inappropriate comments."
    )

# ─── 9. Firebase in §3.7 — Cloud Storage reference ───
replacements["R09-Firebase-37"] = replace_across_doc(
    "Firebase Cloud Storage is reserved for static assets such as user profile pictures.",
    "profile images are managed through external image hosting (ImgBB) with URLs persisted in Firestore."
)
# Try alternate wording
if replacements["R09-Firebase-37"] == 0:
    replacements["R09-Firebase-37"] = replace_across_doc(
        "while Firebase Cloud Storage is reserved for static assets such as user profile pictures.",
        "while profile images are managed through external image hosting (ImgBB) with URLs persisted in Firestore."
    )

# ─── 10. Database schema: Users — add isBanned ───
replacements["R10-UserSchema"] = replace_across_doc(
    "uid (Primary Key), email, name (Full Name), matricNumber, programme, level, role (Student/Admin), isVerified (Boolean), photoURL, createdAt.",
    "uid (Primary Key), email, name (Full Name), matricNumber, programme, level, role (Student/Admin), isVerified (Boolean), isBanned (Boolean), photoURL, createdAt."
)

# ─── 11. Database schema: Videos — update fields ───
replacements["R11-VideoSchema"] = replace_across_doc(
    "videoId (Primary Key), title, description, department, courseCode (Foreign Key), level, topic, userId (Uploader ID), userEmail, thumbnailUrl, videoUrl, status, views (Number), avgRating (Number), isFlagged (Boolean), flaggedReasons (Array), createdAt.",
    "videoId (YouTube ID), title, description, department, courseCode (Foreign Key), level, topic, userId (Uploader ID), userEmail, thumbnailUrl, videoUrl, status (processing/ready/error), stage, statusMessage, views (Number), avgClarity (Number), avgAccuracy (Number), isFlagged (Boolean), flaggedReasons (Array), brokenLink (Boolean), createdAt, completedAt."
)

# ─── 12. Database schema: Ratings — clarity/accuracy ───
replacements["R12-RatingSchema"] = replace_across_doc(
    "Ratings: ratingId, rating (Numerical value), createdAt, videoId (FK), userId (FK).",
    "Ratings: clarity (Numerical value), accuracy (Numerical value), createdAt, userId (FK)."
)

# ─── 13. Sub-collections label ───
replacements["R13-SubCollections"] = replace_across_doc(
    "Rather than a single polymorphic collection, the schema utilizes specialized collections for specific user actions:",
    "Rather than a single polymorphic collection, the schema utilizes specialized sub-collections nested under each video document (videos/{videoId}/) for specific user actions:"
)

# ─── 14. Chapter 4 claims in §4.1.2 — Cloud Storage ───
replacements["R14-Ch4-CloudStorage"] = replace_across_doc(
    "Cloud Storage is utilized specifically for handling non-video assets, such as user profile pictures and static educational thumbnails, ensuring that the platform\u2019s media footprint is distributed efficiently across the Google Cloud ecosystem.",
    "Profile pictures and user avatars are hosted via a third-party image hosting service (ImgBB), with URLs stored in the Firestore user document. This approach avoids the overhead of managing a dedicated file storage bucket for lightweight static assets."
)
if replacements["R14-Ch4-CloudStorage"] == 0:
    replacements["R14-Ch4-CloudStorage"] = replace_across_doc(
        "Cloud Storage is utilized specifically for handling non-video assets, such as user profile pictures and static educational thumbnails, ensuring that the platform's media footprint is distributed efficiently across the Google Cloud ecosystem.",
        "Profile pictures and user avatars are hosted via a third-party image hosting service (ImgBB), with URLs stored in the Firestore user document. This approach avoids the overhead of managing a dedicated file storage bucket for lightweight static assets."
    )

# ─── 15. Chapter 4 claims — caching ───
replacements["R15-Ch4-Caching"] = replace_across_doc(
    "It implements caching and request restructuring",
    "It implements an in-memory response cache with TTL-based expiration and a daily quota tracker"
)

# ─── 16. Microservice caching in §4.5 ───
replacements["R16-Ch4-BackendCache"] = replace_across_doc(
    "the backend implements caching mechanisms for frequently accessed metadata.",
    "the backend implements an in-memory TTL cache for YouTube API responses and a daily quota tracker to monitor API unit consumption."
)

# ─── 17. Request Restructuring in §4.5 ───
replacements["R17-Ch4-RequestRestructure"] = replace_across_doc(
    "Request Restructuring: The custom Express service bundles metadata updates and utilizes optimized polling intervals for the CRON validation scripts, ensuring the platform remains functional throughout the day without hitting API limits.",
    "Quota-Aware Scheduling: The custom Express service batches YouTube API lookups in chunks of 50 IDs, tracks daily quota consumption against the 10,000-unit daily limit, and caches API responses with a one-hour TTL to minimize redundant calls. The hourly CRON validation script checks remaining quota before each batch and gracefully pauses if limits are approached."
)

# ─── 18. Backup claims in §4.3.2 ───
replacements["R18-Ch4-Backups"] = replace_across_doc(
    "automated backup routines were configured within the Google Cloud Console to run on a scheduled basis",
    "an automated backup routine runs daily at 2:00 AM UTC via a scheduled CRON job in the Node.js backend, capturing document counts and collection snapshots from Firestore and persisting backup records for administrative audit"
)

# ─── 19. Backup claims continued ───
replacements["R19-Ch4-BackupRestore"] = replace_across_doc(
    "This ensures that even in the unlikely event of data corruption or accidental deletion, the platform\u2019s content catalog, user profiles, and interaction records can be restored without significant loss.",
    "This ensures that the platform maintains a verifiable record of data integrity, with backup logs accessible through the administrator dashboard for audit purposes."
)
if replacements["R19-Ch4-BackupRestore"] == 0:
    replacements["R19-Ch4-BackupRestore"] = replace_across_doc(
        "This ensures that even in the unlikely event of data corruption or accidental deletion, the platform's content catalog, user profiles, and interaction records can be restored without significant loss.",
        "This ensures that the platform maintains a verifiable record of data integrity, with backup logs accessible through the administrator dashboard for audit purposes."
    )

# ─── 20. XSS in §4.2.4 ───
replacements["R20-Ch4-XSS"] = replace_across_doc(
    "Input Sanitization: Ensured that comment sections were protected against common vulnerabilities like Cross-Site Scripting (XSS).",
    "Input Sanitization: Comment sections are protected against Cross-Site Scripting (XSS) through Vue.js\u2019s default template rendering, which automatically escapes HTML entities in all user-generated content rendered via mustache interpolation ({{ }}), preventing injection of malicious scripts."
)
if replacements["R20-Ch4-XSS"] == 0:
    replacements["R20-Ch4-XSS"] = replace_across_doc(
        "Input Sanitization: Ensured that comment sections were protected against common vulnerabilities like Cross-Site Scripting (XSS).",
        "Input Sanitization: Comment sections are protected against Cross-Site Scripting (XSS) through Vue.js's default template rendering, which automatically escapes HTML entities in all user-generated content rendered via mustache interpolation ({{ }}), preventing injection of malicious scripts."
    )

# ─── 21. Usability in §4.2.3 — Bootstrap 5 mention ───
replacements["R21-Ch4-Responsive"] = replace_across_doc(
    "The interface was tested across smartphones and laptops using Bootstrap 5, ensuring the 3-click rule was maintained for easy navigation.",
    "The interface was tested across smartphones and laptops using responsive CSS breakpoints, ensuring the 3-click rule was maintained for easy navigation."
)

# ─── 22. Comment via "real-time WebSocket sync" (§4.2.3 test table or §4.3 test case) ───
replacements["R22-WebSocket"] = replace_across_doc(
    "real-time WebSocket sync",
    "real-time Firestore sync"
)

# ─── 23. Admin flags wording in §7 Governance ───
replacements["R23-AdminFlags"] = replace_across_doc(
    "The system shall allow administrators to flag videos that violate content guidelines or are deemed inappropriate.",
    "The system shall allow authenticated users to flag videos that violate content guidelines, with administrators reviewing and acting on flagged content through a dedicated moderation queue."
)

# ─── 24. Admin reports wording ───
replacements["R24-AdminReports"] = replace_across_doc(
    "The system shall allow administrators to view reports on flagged content, user activity, and platform usage.",
    "The system shall allow administrators to view aggregate platform analytics and review flagged content through dedicated dashboard panels."
)

# ─── 25. Capability 43 — mobile-first with Bootstrap 5 ───
replacements["R25-Cap43"] = replace_across_doc(
    "The system shall provide an intuitive, mobile-first responsive interface built with Bootstrap 5 that functions seamlessly on smartphones, tablets, and desktop PCs.",
    "The system shall provide an intuitive, responsive interface built with custom CSS3 and Bootstrap 5 utilities that functions seamlessly on smartphones, tablets, and desktop PCs."
)

# ─── 26. Search by keywords cap 17 ───
replacements["R26-Cap17"] = replace_across_doc(
    "The system shall allow users to search for videos by course code, department, level, topic, and academic keywords.",
    "The system shall allow users to search for videos by title, topic, and course code via text search, with dropdown filters available for Department, Course Code, and Level."
)

# ─── 27. Cap 19 — Topic filter ───
replacements["R27-Cap19"] = replace_across_doc(
    "The system shall provide filtering options to narrow results by Department, Course Code, Level, and Topic.",
    "The system shall provide filtering options to narrow results by Department, Course Code, and Level, with Topic discoverable via text search."
)

# ─── 28. Cap 20 — ranked relevant ───
replacements["R28-Cap20"] = replace_across_doc(
    "The system shall display search results in a ranked, relevant order based on the user\u2019s query.",
    "The system shall display search results sorted by recency or popularity (view count), with text matching applied to titles, topics, and course codes."
)
if replacements["R28-Cap20"] == 0:
    replacements["R28-Cap20"] = replace_across_doc(
        "The system shall display search results in a ranked, relevant order based on the user's query.",
        "The system shall display search results sorted by recency or popularity (view count), with text matching applied to titles, topics, and course codes."
    )

# ══════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════
doc.save(DOC_PATH)

# Report
print("=" * 60)
print("DOCX UPDATE REPORT")
print("=" * 60)
applied = {k: v for k, v in replacements.items() if v > 0}
missed = {k: v for k, v in replacements.items() if v == 0}

print(f"\nApplied: {len(applied)} / {len(replacements)}")
for k, v in applied.items():
    print(f"  [OK] {k} ({v} replacement(s))")

if missed:
    print(f"\nMissed: {len(missed)} (text not found exactly — may need manual check)")
    for k in missed:
        print(f"  [!!] {k}")

print(f"\nFile saved to: {DOC_PATH}")
