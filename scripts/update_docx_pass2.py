"""
Update CHAPTER THREE (Updated).docx — second pass targeting the ACTUAL paragraphs in the docx.
"""
from docx import Document

DOC_PATH = r"c:\Projects\eduvid\CHAPTER THREE (Updated).docx"
doc = Document(DOC_PATH)

def replace_in_paragraph(paragraph, old_text, new_text):
    full = paragraph.text
    if old_text not in full:
        return False
    new_full = full.replace(old_text, new_text)
    if paragraph.runs:
        paragraph.runs[0].text = new_full
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_full
    return True

def replace_across_doc(old_text, new_text, label=""):
    count = 0
    for p in doc.paragraphs:
        if replace_in_paragraph(p, old_text, new_text):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if replace_in_paragraph(p, old_text, new_text):
                        count += 1
    status = "[OK]" if count > 0 else "[!!]"
    print(f"  {status} {label}: {count} replacement(s)")
    return count

print("=" * 60)
print("DOCX UPDATE - PASS 2 (actual paragraph text)")
print("=" * 60)

# ─── R03: Bootstrap 5 in §3.7 (actual docx text) ───
replace_across_doc(
    "Bootstrap 5 (UI Framework): To ensure accessibility across devices, Bootstrap 5 provides the responsive grid system and pre-designed UI components",
    "Custom CSS3 with Bootstrap 5 Utilities (UI Styling): The platform\u2019s visual identity is built on a custom CSS3 design system featuring CSS Grid layouts, Flexbox positioning, and Glassmorphism effects achieved via backdrop-filter. Bootstrap 5 is included as a supplementary utility library, with its icon set (Bootstrap Icons) used extensively throughout the interface",
    "R03-Bootstrap-37"
)

# ─── R05: Cloud Storage — actual docx doesn't have the exact §3.4.2 text from the pasted doc ───
# The docx has a different server section. Let's check what's there:
# para 71: "Firebase Hosting: The entire application is deployed via Firebase Hosting..."
# The docx doesn't mention "Cloud Storage" at all in §3.4.2. 
# So R05 is N/A for the docx (it was already different).

# ─── R06: API Quota Management — the docx doesn't have this section ───
# The original docx doesn't have the Node.js/Express/Render backend section.
# R06 is N/A.

# ─── R10-R13: Database Schema — the docx has completely different schema ───
# Replace the old schema paragraphs with updated ones

# Users collection (para 96-97)
replace_across_doc(
    "Fields: Fields: uid (Google/Manual ID), firstName, lastName, email, matricNumber, role (Student/Admin)",
    "Fields: uid (Primary Key), email, name (Full Name), matricNumber, programme, level, role (Student/Admin), isVerified (Boolean), isBanned (Boolean), photoURL, createdAt",
    "R10-UserSchema"
)

# Videos collection (para 98-99)
replace_across_doc(
    "Fields: videoId (YouTube ID), courseCode (FK), uploaderId (FK), isFlagged (Boolean), status (active/deleted)",
    "Fields: videoId (YouTube ID), title, description, department, courseCode (FK), level, topic, userId (Uploader ID), userEmail, thumbnailUrl, videoUrl, status (processing/ready/error), stage, statusMessage, views (Number), avgClarity (Number), avgAccuracy (Number), isFlagged (Boolean), flaggedReasons (Array), brokenLink (Boolean), createdAt, completedAt",
    "R11-VideoSchema"
)

# Courses collection (para 100-101)
replace_across_doc(
    "Fields: courseCode (Primary Key, e.g., SEN401), courseTitle.",
    "Fields: code (e.g., SEN401), department, videoCount, createdAt.",
    "R12-CourseSchema"
)

# Interactions collection header (para 102)
replace_across_doc(
    "4. Interactions Collection (Comments & Ratings)",
    "4. Interaction Sub-Collections (nested under videos/{videoId}/)",
    "R13a-SubCollLabel"
)

# Interactions description (para 103)
replace_across_doc(
    "A polymorphic collection for peer feedback and moderation.",
    "Specialized sub-collections for specific user actions:",
    "R13b-SubCollDesc"
)

# Interactions fields (para 104)
replace_across_doc(
    "Fields: interactionId, videoId (FK), userId (FK), interactionType (comment/flag), content, createdAt",
    "Comments: text, userName, userEmail, createdAt, userId. Ratings: clarity (Numerical value), accuracy (Numerical value), createdAt, userId. Flags: reason, userEmail, createdAt, userId",
    "R13c-SubCollFields"
)

# ─── R14-R22: Chapter 4 content is NOT in this docx ───
# The chapter_three.docx doesn't contain Chapter 4 content.
# Those replacements are N/A — Chapter 4 was already updated in chapter_four.md

# ─── R23-R28: System capabilities text is NOT in this docx ───
# These are in system_capabilities.md which was already updated.

# ─── Firebase Hosting reference — update to mention Vercel/Render ───
replace_across_doc(
    "Firebase Hosting: The entire application is deployed via Firebase Hosting, which provides global delivery through a Content Delivery Network (CDN)",
    "Vercel & Render Hosting: The application utilizes a hybrid hosting strategy. The Vue.js frontend is deployed on Vercel, leveraging an edge network for fast global delivery and SSL encryption. The Node.js backend is hosted on Render, providing a managed environment for API logic and background synchronization tasks",
    "R-FirebaseHosting-342"
)

replace_across_doc(
    "Firebase Hosting: The application is deployed using Firebase Hosting, which provides a production-grade, SSL-encrypted environment",
    "Vercel (Frontend Hosting): Used to deploy the Vue.js frontend, leveraging an edge network for fast global delivery and continuous deployment from GitHub. Render (Backend Hosting): Provides a managed environment for the Node.js/Express API logic and background synchronization tasks",
    "R-FirebaseHosting-37"
)

# ─── Add missing tools to §3.7: Vite, Pinia, Node.js/Express ───
# The docx §3.7 only lists Vue.js, Bootstrap 5, Firebase, Firebase Hosting, GitHub
# Missing: Vite, Pinia, Node.js/Express, YouTube Data API v3

# Insert after Vue.js paragraph by finding it
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("Vue.js 3 (Frontend Framework):"):
        # Add Vite after Vue.js
        # We'll append to the Vue.js paragraph text
        replace_in_paragraph(p, 
            "Vue.js is utilized as the primary frontend framework to build a Single Page Application (SPA). Its reactive component system allows for a seamless user experience, where the interface updates dynamically without requiring the browser to reload pages entirely.",
            "Vue.js is utilized as the primary frontend framework to build a Single Page Application (SPA) using the Composition API. Its reactive component system allows for a seamless user experience, where the interface updates dynamically without requiring the browser to reload pages entirely. Vite serves as the build tool and development server, providing fast Hot Module Replacement (HMR) and optimized bundling for production. Pinia is employed as the official state management library for Vue 3 to manage global application state, including user authentication and real-time upload progress."
        )
        break

# Add Node.js/Express/YouTube API info to Firebase paragraph
replace_across_doc(
    "Firebase provides the comprehensive backend infrastructure, eliminating the need to manage physical servers",
    "Firebase provides the comprehensive backend infrastructure, eliminating the need to manage physical servers. A custom Node.js & Express microservice is hosted on Render to handle specialized logic such as the YouTube validation pipeline, CRON-style routines for link moderation, API quota management, and automated backups. The YouTube Data API v3 is utilized for managing academic video metadata and leveraging YouTube\u2019s global infrastructure for hosting and transcoding large video files",
    "R-AddNodeExpress"
)

doc.save(DOC_PATH)

print(f"\nFile saved to: {DOC_PATH}")
