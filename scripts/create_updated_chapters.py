import os
from docx import Document
from copy import deepcopy

# Find the main docx dynamically to avoid unicode filename issues
_dir = r"c:\Projects\eduvid"
_candidates = [f for f in os.listdir(_dir) if f.startswith("DESIGN AND") and f.endswith(".docx")]
if not _candidates:
    raise FileNotFoundError("Could not find the main project docx")
SRC = os.path.join(_dir, _candidates[0])
CH3_OUT = r"c:\Projects\eduvid\CHAPTER THREE (Updated).docx"
CH4_OUT = r"c:\Projects\eduvid\CHAPTER FOUR (Updated).docx"

src_doc = Document(SRC)

# ── Find boundaries ──
ch3_start = ch4_start = ch5_start = None
for i, p in enumerate(src_doc.paragraphs):
    txt = p.text.strip()
    if txt == "CHAPTER THREE": ch3_start = i
    elif txt == "CHAPTER FOUR": ch4_start = i
    elif txt == "CHAPTER FIVE": ch5_start = i

print(f"Ch3: {ch3_start}-{ch4_start-1}, Ch4: {ch4_start}-{ch5_start-1}")

# ══════════════════════════════════════════
# Helper to copy paragraphs into a new doc
# ══════════════════════════════════════════
def copy_paragraph(src_para, dst_doc):
    """Copy a paragraph preserving style and run formatting."""
    new_para = dst_doc.add_paragraph()
    new_para.style = src_para.style
    # Copy paragraph formatting
    if src_para.paragraph_format.alignment is not None:
        new_para.paragraph_format.alignment = src_para.paragraph_format.alignment
    # Copy runs
    for run in src_para.runs:
        new_run = new_para.add_run(run.text)
        new_run.bold = run.bold
        new_run.italic = run.italic
        new_run.underline = run.underline
        if run.font.size:
            new_run.font.size = run.font.size
        if run.font.name:
            new_run.font.name = run.font.name
        if run.font.color and run.font.color.rgb:
            new_run.font.color.rgb = run.font.color.rgb
    # If no runs (plain text paragraph)
    if not src_para.runs and src_para.text:
        new_para.text = src_para.text
    return new_para

def extract_chapter(start, end, output_path):
    """Extract paragraphs [start..end] into a new docx."""
    new_doc = Document()
    # Remove the default empty paragraph
    if new_doc.paragraphs:
        p = new_doc.paragraphs[0]._element
        p.getparent().remove(p)
    for i in range(start, end + 1):
        copy_paragraph(src_doc.paragraphs[i], new_doc)
    new_doc.save(output_path)
    return new_doc

# ── Extract ──
print("Extracting Chapter 3...")
extract_chapter(ch3_start, ch4_start - 1, CH3_OUT)
print(f"  Saved to {CH3_OUT}")

print("Extracting Chapter 4...")
extract_chapter(ch4_start, ch5_start - 1, CH4_OUT)
print(f"  Saved to {CH4_OUT}")

# ══════════════════════════════════════════
# Now apply corrections to each file
# ══════════════════════════════════════════
def replace_in_doc(doc, old_text, new_text, label=""):
    count = 0
    for p in doc.paragraphs:
        if old_text in p.text:
            new_full = p.text.replace(old_text, new_text)
            if p.runs:
                p.runs[0].text = new_full
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = new_full
            count += 1
    status = "[OK]" if count > 0 else "[!!]"
    print(f"  {status} {label}: {count}")
    return count

# ────────────────────────────────────────
# CHAPTER 3 CORRECTIONS
# ────────────────────────────────────────
print("\n" + "=" * 60)
print("APPLYING CHAPTER 3 CORRECTIONS")
print("=" * 60)

ch3 = Document(CH3_OUT)

# 1. Bootstrap in §3.3.2 NFR
replace_in_doc(ch3,
    "A mobile-first design using Bootstrap 5 to ensure the platform works seamlessly on smartphones, tablets, and PCs.",
    "A responsive interface using custom CSS Grid and Flexbox layouts with Bootstrap 5 utility support, ensuring the platform works seamlessly on smartphones, tablets, and PCs.",
    "NFR-Responsiveness")

# 2. Bootstrap in §3.4.1
replace_in_doc(ch3,
    "Bootstrap 5: To ensure a mobile-first experience for Dominion University students, Bootstrap 5 provides a responsive grid system. This ensures the platform is accessible across various devices, including smartphones, tablets, and laptops",
    "Custom CSS3 with Bootstrap 5 Utilities: The platform\u2019s responsive design is achieved through a custom CSS3 design system built with CSS Grid and Flexbox layouts, featuring Glassmorphism visual effects via backdrop-filter. Bootstrap 5 is included as a utility library for supplementary spacing and icon support, while all core layout and component styling is handled through purpose-built CSS classes and scoped @media breakpoints to ensure the interface adapts across smartphones, tablets, and laptops",
    "Bootstrap-341")

# 3. Bootstrap in §3.7
replace_in_doc(ch3,
    "Bootstrap 5 & CSS3 (UI Framework): Provides a responsive grid system and pre-styled components. This ensures the platform adapts to various screen sizes, from smartphones to laptops, while custom CSS3 handles advanced effects like Glassmorphism.",
    "Custom CSS3 with Bootstrap 5 Utilities (UI Styling): The platform\u2019s visual identity is built on a custom CSS3 design system featuring CSS Grid layouts, Flexbox positioning, and Glassmorphism effects achieved via backdrop-filter. Bootstrap 5 is included as a supplementary utility library, with its icon set (Bootstrap Icons) used extensively throughout the interface. All responsive breakpoints are handled through custom @media queries.",
    "Bootstrap-37a")

# Also try the alternate docx wording
replace_in_doc(ch3,
    "Bootstrap 5 (UI Framework): To ensure accessibility across devices, Bootstrap 5 provides the responsive grid system and pre-styled components",
    "Custom CSS3 with Bootstrap 5 Utilities (UI Styling): The platform\u2019s visual identity is built on a custom CSS3 design system featuring CSS Grid layouts, Flexbox positioning, and Glassmorphism effects. Bootstrap 5 is included as a supplementary utility library with its icon set used extensively",
    "Bootstrap-37b")

# 4. YouTube IFrame API
replace_in_doc(ch3,
    "uses the YouTube IFrame API to embed the video players",
    "renders embedded YouTube video players using standard iframe embeds with YouTube\u2019s /embed/ endpoint",
    "IFrame-API")
replace_in_doc(ch3,
    "YouTube IFrame API to embed the video players directly into the student's dashboard",
    "standard iframe embeds with YouTube\u2019s /embed/ endpoint to render video players directly into the student's dashboard",
    "IFrame-API-alt")

# 5. Cloud Storage
replace_in_doc(ch3,
    "Cloud Storage: Reserved for storing static assets, such as user profile pictures and application icons.",
    "Image Hosting: Profile pictures and user avatars are hosted via a third-party image hosting service (ImgBB), with the resulting URLs stored directly in the Firestore user document.",
    "CloudStorage-342")

# 6. API Quota Management
replace_in_doc(ch3,
    "API Quota Management: Logic to cache responses and restructure requests to stay within the YouTube Data API v3 daily limits.",
    "API Quota Management: Logic to track daily API usage against the YouTube Data API v3 quota limits. The backend implements an in-memory response cache with TTL-based expiration and a quota tracker that monitors daily unit consumption, ensuring that automated validation routines stay within the daily allowance.",
    "QuotaMgmt-342")

# 7. Search & Discovery
replace_in_doc(ch3,
    "Search and Discovery Engine: Filtering and search functionality based on course codes and academic keywords.",
    "Search and Discovery Engine: Text-based search across video titles, topics, and course codes, with dropdown filtering by Department, Course Code, and Level.",
    "SearchDiscovery")

# 8. Admin flagging
replace_in_doc(ch3,
    'moderate content by \u201cflagging\u201d videos and managing user permissions.',
    "review flagged content, manage user permissions, and moderate inappropriate comments.",
    "AdminGov-smart")
replace_in_doc(ch3,
    'moderate content by "flagging" videos and managing user permissions.',
    "review flagged content, manage user permissions, and moderate inappropriate comments.",
    "AdminGov-straight")

# 9. Firebase Cloud Storage in §3.7
replace_in_doc(ch3,
    "Firebase Cloud Storage is reserved for static assets such as user profile pictures",
    "profile images are managed through external image hosting (ImgBB) with URLs persisted in Firestore",
    "Firebase-37-storage")
replace_in_doc(ch3,
    "while Firebase Cloud Storage is reserved for static assets such as user profile pictures",
    "while profile images are managed through external image hosting (ImgBB) with URLs persisted in Firestore",
    "Firebase-37-storage-alt")

# 10. Database Schema: Users — add isBanned
replace_in_doc(ch3,
    "uid (Primary Key), email, name (Full Name), matricNumber, programme, level, role (Student/Admin), isVerified (Boolean), photoURL, createdAt.",
    "uid (Primary Key), email, name (Full Name), matricNumber, programme, level, role (Student/Admin), isVerified (Boolean), isBanned (Boolean), photoURL, createdAt.",
    "UserSchema")

# 11. Database Schema: Videos
replace_in_doc(ch3,
    "videoId (Primary Key), title, description, department, courseCode (Foreign Key), level, topic, userId (Uploader ID), userEmail, thumbnailUrl, videoUrl, status, views (Number), avgRating (Number), isFlagged (Boolean), flaggedReasons (Array), createdAt.",
    "videoId (YouTube ID), title, description, department, courseCode (FK), level, topic, userId (Uploader ID), userEmail, thumbnailUrl, videoUrl, status (processing/ready/error), stage, statusMessage, views (Number), avgClarity (Number), avgAccuracy (Number), isFlagged (Boolean), flaggedReasons (Array), brokenLink (Boolean), createdAt, completedAt.",
    "VideoSchema")

# 12. Ratings schema
replace_in_doc(ch3,
    "Ratings: ratingId, rating (Numerical value), createdAt, videoId (FK), userId (FK).",
    "Ratings: clarity (Numerical value), accuracy (Numerical value), createdAt, userId.",
    "RatingSchema")

# 13. Sub-collections label
replace_in_doc(ch3,
    "Rather than a single polymorphic collection, the schema utilizes specialized collections for specific user actions:",
    "Rather than a single polymorphic collection, the schema utilizes specialized sub-collections nested under each video document (videos/{videoId}/) for specific user actions:",
    "SubCollections")

# 14. Firebase Hosting -> Vercel/Render
replace_in_doc(ch3,
    "Firebase Hosting: The entire application is deployed via Firebase Hosting",
    "Vercel & Render Hosting: The application utilizes a hybrid hosting strategy. The Vue.js frontend is deployed on Vercel",
    "FirebaseHosting-342")
replace_in_doc(ch3,
    "which provides global delivery through a Content Delivery Network (CDN)",
    "leveraging an edge network for fast global delivery and SSL encryption. The Node.js backend is hosted on Render, providing a managed environment for API logic and background synchronization tasks",
    "FirebaseHosting-342-cdn")

replace_in_doc(ch3,
    "Firebase Hosting: The application is deployed using Firebase Hosting",
    "Vercel (Frontend Hosting): Used to deploy the Vue.js frontend, leveraging an edge network for fast global delivery. Render (Backend Hosting): Provides a managed environment for the Node.js/Express API logic",
    "FirebaseHosting-37")

# 15. Add Node.js/Express/YouTube API info
replace_in_doc(ch3,
    "Firebase provides the comprehensive backend infrastructure, eliminating the need to manage physical servers",
    "Firebase provides the comprehensive backend infrastructure, eliminating the need to manage physical servers. A custom Node.js & Express microservice is hosted on Render to handle specialized logic such as the YouTube validation pipeline, CRON-style routines for link moderation, API quota management with in-memory caching, and automated daily backups. The YouTube Data API v3 is utilized for managing academic video metadata and leveraging YouTube\u2019s global infrastructure for hosting and transcoding",
    "AddNodeExpress")

# 16. Old schema from the original docx
replace_in_doc(ch3,
    "Fields: Fields: uid (Google/Manual ID), firstName, lastName, email, matricNumber, role (Student/Admin)",
    "Fields: uid (Primary Key), email, name (Full Name), matricNumber, programme, level, role (Student/Admin), isVerified (Boolean), isBanned (Boolean), photoURL, createdAt",
    "OldUserSchema")

replace_in_doc(ch3,
    "Fields: videoId (YouTube ID), courseCode (FK), uploaderId (FK), isFlagged (Boolean), status (active/deleted)",
    "Fields: videoId (YouTube ID), title, description, department, courseCode (FK), level, topic, userId (Uploader ID), userEmail, thumbnailUrl, videoUrl, status (processing/ready/error), stage, views (Number), avgClarity (Number), avgAccuracy (Number), isFlagged (Boolean), flaggedReasons (Array), brokenLink (Boolean), createdAt, completedAt",
    "OldVideoSchema")

replace_in_doc(ch3,
    "Fields: courseCode (Primary Key, e.g., SEN401), courseTitle.",
    "Fields: code (e.g., SEN401), department, videoCount, createdAt.",
    "OldCourseSchema")

replace_in_doc(ch3,
    "4. Interactions Collection (Comments & Ratings)",
    "4. Interaction Sub-Collections (nested under videos/{videoId}/)",
    "OldInterLabel")

replace_in_doc(ch3,
    "A polymorphic collection for peer feedback and moderation.",
    "Specialized sub-collections for specific user actions:",
    "OldInterDesc")

replace_in_doc(ch3,
    "Fields: interactionId, videoId (FK), userId (FK), interactionType (comment/flag), content, createdAt",
    "Comments: text, userName, userEmail, createdAt, userId. Ratings: clarity (Numerical), accuracy (Numerical), createdAt, userId. Flags: reason, userEmail, createdAt, userId",
    "OldInterFields")

ch3.save(CH3_OUT)
print(f"\nChapter 3 saved: {CH3_OUT}")

# ────────────────────────────────────────
# CHAPTER 4 CORRECTIONS
# ────────────────────────────────────────
print("\n" + "=" * 60)
print("APPLYING CHAPTER 4 CORRECTIONS")
print("=" * 60)

ch4 = Document(CH4_OUT)

# 1. Backups in §4.3.2
replace_in_doc(ch4,
    "automated backup routines were configured within the Google Cloud Console to run on a scheduled basis",
    "an automated backup routine runs daily at 2:00 AM UTC via a scheduled CRON job in the Node.js backend, capturing document counts and collection snapshots from Firestore and persisting backup records for administrative audit",
    "Backups-432")

# 2. Backup restore claim
replace_in_doc(ch4,
    "This ensures that even in the unlikely event of data corruption or accidental deletion, the platform's content catalog, user profiles, and interaction records can be restored without significant loss.",
    "This ensures that the platform maintains a verifiable record of data integrity, with backup logs accessible through the administrator dashboard for audit purposes.",
    "BackupRestore")
replace_in_doc(ch4,
    "This ensures that even in the unlikely event of data corruption or accidental deletion, the platform\u2019s content catalog, user profiles, and interaction records can be restored without significant loss.",
    "This ensures that the platform maintains a verifiable record of data integrity, with backup logs accessible through the administrator dashboard for audit purposes.",
    "BackupRestore-smart")

# 3. XSS in §4.2.4
replace_in_doc(ch4,
    "Input Sanitization: Ensured that comment sections were protected against common vulnerabilities like Cross-Site Scripting (XSS).",
    "Input Sanitization: Comment sections are protected against Cross-Site Scripting (XSS) through Vue.js\u2019s default template rendering, which automatically escapes HTML entities in all user-generated content rendered via mustache interpolation, preventing injection of malicious scripts.",
    "XSS")

# 4. Bootstrap 5 in §4.2.3
replace_in_doc(ch4,
    "The interface was tested across smartphones and laptops using Bootstrap 5, ensuring the 3-click rule was maintained",
    "The interface was tested across smartphones and laptops using responsive CSS breakpoints, ensuring the 3-click rule was maintained",
    "Responsive-423")

# 5. WebSocket
replace_in_doc(ch4,
    "real-time WebSocket sync",
    "real-time Firestore sync",
    "WebSocket")

# 6. Caching claims in §4.5
replace_in_doc(ch4,
    "the backend implements caching mechanisms for frequently accessed metadata",
    "the backend implements an in-memory TTL cache for YouTube API responses and a daily quota tracker to monitor API unit consumption",
    "BackendCache")

# 7. Request Restructuring in §4.5
replace_in_doc(ch4,
    "Request Restructuring: The custom Express service bundles metadata updates and utilizes optimized polling intervals for the CRON validation scripts, ensuring the platform remains functional throughout the day without hitting API limits.",
    "Quota-Aware Scheduling: The custom Express service batches YouTube API lookups in chunks of 50 IDs, tracks daily quota consumption against the 10,000-unit daily limit, and caches API responses with a one-hour TTL to minimize redundant calls. The hourly CRON validation script checks remaining quota before each batch and gracefully pauses if limits are approached.",
    "RequestRestructure")

# 8. Maintenance approach backup mention
replace_in_doc(ch4,
    "Automated backups were configured within the Google Cloud Console on a scheduled basis, covering Firestore data to protect against accidental deletion or corruption.",
    "Automated backups run daily at 2:00 AM UTC via a CRON job in the Node.js backend, capturing collection-level document counts from Firestore and persisting backup records for administrative review.",
    "MaintenanceBackups")

# 9. Cloud Storage in §4.1.2
replace_in_doc(ch4,
    "Cloud Storage is utilized specifically for handling non-video assets, such as user profile pictures and static educational thumbnails",
    "Profile pictures and user avatars are hosted via a third-party image hosting service (ImgBB), with URLs stored in the Firestore user document",
    "CloudStorage-Ch4")
replace_in_doc(ch4,
    "ensuring that the platform\u2019s media footprint is distributed efficiently across the Google Cloud ecosystem",
    "This approach avoids the overhead of managing a dedicated file storage bucket for lightweight static assets",
    "CloudStorage-Ch4-tail-smart")
replace_in_doc(ch4,
    "ensuring that the platform's media footprint is distributed efficiently across the Google Cloud ecosystem",
    "This approach avoids the overhead of managing a dedicated file storage bucket for lightweight static assets",
    "CloudStorage-Ch4-tail")

# 10. Caching claim in §4.1.2
replace_in_doc(ch4,
    "It implements caching and request restructuring",
    "It implements an in-memory response cache with TTL-based expiration and a daily quota tracker",
    "Caching-412")

ch4.save(CH4_OUT)
print(f"\nChapter 4 saved: {CH4_OUT}")
print("\nDone!")
