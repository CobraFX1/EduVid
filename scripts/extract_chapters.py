"""Find chapter boundaries and extract full text for chapters 3 and 4."""
from docx import Document

DOC_PATH = r"c:\Projects\eduvid\DESIGN AND IMPLEMENTATION OF A WEB – BASED EDUCATIONAL VIDEO PLATFORM FOR PEER-TO-PEER LEARNING.docx"
doc = Document(DOC_PATH)

# Find chapter boundaries
ch3_start = ch4_start = ch5_start = None
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt == "CHAPTER THREE" or txt.startswith("CHAPTER THREE"):
        ch3_start = i
        print(f"Chapter 3 starts at paragraph {i}: {txt}")
    elif txt == "CHAPTER FOUR" or txt.startswith("CHAPTER FOUR"):
        ch4_start = i
        print(f"Chapter 4 starts at paragraph {i}: {txt}")
    elif txt == "CHAPTER FIVE" or txt.startswith("CHAPTER FIVE"):
        ch5_start = i
        print(f"Chapter 5 starts at paragraph {i}: {txt}")

print(f"\nChapter 3: paragraphs {ch3_start} to {ch4_start - 1}")
print(f"Chapter 4: paragraphs {ch4_start} to {ch5_start - 1}")

# Dump full text of chapter 3
print("\n" + "=" * 60)
print("CHAPTER 3 FULL TEXT")
print("=" * 60)
for i in range(ch3_start, ch4_start):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"[{i:03d}] {p.text}")

print("\n" + "=" * 60)
print("CHAPTER 4 FULL TEXT")
print("=" * 60)
for i in range(ch4_start, ch5_start):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"[{i:03d}] {p.text}")
